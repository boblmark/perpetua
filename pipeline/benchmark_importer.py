#!/usr/bin/env python3
"""Arena Top-10 标杆账号 benchmark importer。

读 xlsx 中的 Top_10_AI_Comedy_Benchmark sheet，存为研究数据
/home/hatch/data/perpetua/research/benchmark_YYYYMMDD.json。
注意：这是竞品情报，不进 pipeline_queue（那是咱们自己的选题队列）。

用法：
  python3 benchmark_importer.py --ingest-drops     # 处理 drops/ 下含 benchmark sheet 的文件
  python3 benchmark_importer.py --ingest <xlsx>
  python3 benchmark_importer.py --dry-run <xlsx>
"""
sys.path.insert(0, "/home/hatch/data/perpetua")
from json_safe import atomic_write_json
import json
import os
import re
import shutil
import sys
from datetime import datetime, timezone

DATA = "/home/hatch/data"
DROPS = os.path.join(DATA, "perpetua", "arena_agent", "drops")
ARCHIVE = os.path.join(DATA, "perpetua", "arena_agent", "archive")
RESEARCH = os.path.join(DATA, "perpetua", "research")
SHEET = "Top_10_AI_Comedy_Benchmark"


def parse_benchmark(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        return None, f"无 {SHEET} sheet（仅有：{wb.sheetnames}）"
    ws = wb[SHEET]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return None, "sheet 为空"
    headers = [(str(c or "").strip()) for c in rows[0]]
    accts = []
    for r in rows[1:]:
        if not any(r):
            continue
        d = dict(zip(headers, [str(c).strip() if c is not None else "" for c in r]))
        if not d.get(headers[0]):
            continue
        accts.append(d)
    return {"headers": headers, "accounts": accts,
            "source_file": os.path.basename(path),
            "imported_at": datetime.now(timezone.utc).isoformat()}, None


def ingest(path, dry_run=False):
    data, err = parse_benchmark(path)
    if err:
        print(f"SKIP {os.path.basename(path)}：{err}")
        return False
    print(f"解析 {len(data['accounts'])} 个标杆账号 ← {os.path.basename(path)}")
    for a in data["accounts"]:
        name = a.get(data["headers"][0], "")
        print(f"  - {name[:50]}")
    if dry_run:
        print("dry-run，不落盘")
        return True
    os.makedirs(RESEARCH, exist_ok=True)
    day = datetime.now(timezone.utc).strftime("%Y%m%d")
    out = os.path.join(RESEARCH, f"benchmark_{day}.json")
    # 同一天多次导入则追加版本号
    if os.path.exists(out):
        out = os.path.join(RESEARCH, f"benchmark_{day}_{datetime.now(timezone.utc).strftime('%H%M%S')}.json")
    atomic_write_json(out, data)
    print(f"落盘 → {out}")
    dst = os.path.join(ARCHIVE, os.path.basename(path))
    os.makedirs(ARCHIVE, exist_ok=True)
    shutil.move(path, dst)
    print(f"归档 → {dst}")
    return True


def main():
    os.makedirs(DROPS, exist_ok=True)
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    cmd, arg = sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else ""
    if cmd == "--ingest-drops":
        files = sorted(f for f in os.listdir(DROPS) if f.endswith(".xlsx"))
        if not files:
            print("drops/ 为空")
            return
        for f in files:
            ingest(os.path.join(DROPS, f))
    elif cmd in ("--ingest", "--dry-run"):
        if not arg:
            print("需要文件路径")
            sys.exit(1)
        ingest(arg, dry_run=(cmd == "--dry-run"))
    else:
        print(__doc__)
        sys.exit(1)


if __name__ == "__main__":
    main()
