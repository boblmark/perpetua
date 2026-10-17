# -*- coding: utf-8 -*-
"""生成 样片库_20261014.xlsx —— 从 288 条真实数据中精选的可复核样片索引。"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from kb_common import (ROWS, PLAT_CN, build_samples, sample_reason_auto, num, fmt, is_viral)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
OUT_DIR = os.environ.get("PIPELINE_OUT", os.path.dirname(os.path.dirname(os.path.abspath(__file__))))


OUT = os.path.join(OUT_DIR, "样片库_20261014.xlsx")

HEAD = ["序号", "平台", "账号", "账号粉丝", "视频链接", "标题", "封面文字(OCR)", "播放", "点赞", "评论",
        "转发", "收藏", "弹幕", "时长(秒)", "状态", "选题分类", "开场钩子", "量化特征（由本行真实数字计算）", "入选理由"]

WIDTH = [5, 8, 16, 11, 40, 46, 22, 11, 10, 9, 9, 9, 8, 9, 7, 15, 15, 46, 40]

THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def main():
    S = build_samples()
    wb = Workbook()
    ws = wb.active
    ws.title = "样片库"

    # 说明区
    notes = [
        "样片库 · 精选自 competitor_viral_20261014.xlsx（288 条真实数据，5 个平台）",
        "选取原则：每平台覆盖头部爆款 + 对照组中位样本 + 结构特殊样本（超长/超短/高收藏/高转发），全部含可点击原视频链接。",
        "「量化特征」列由该行真实公开数字直接计算（例：点赞 325194 ÷ 播放 8327176 = 3.91%），可自行复核。",
        "「封面文字」为封面图 OCR 实测结果，未识别不等于封面无文字。",
        "本表不含任何估算值；N/A 表示平台未公开该字段。抖音不公开播放量，故按点赞排序。",
    ]
    for i, t in enumerate(notes, 1):
        ws.cell(row=i, column=1, value=t)
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=len(HEAD))
        c = ws.cell(row=i, column=1)
        c.font = Font(size=11 if i == 1 else 9, bold=(i == 1), color="1F3864" if i == 1 else "595959")
        c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 22

    hr = len(notes) + 2
    for j, h in enumerate(HEAD, 1):
        c = ws.cell(row=hr, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor="2F5597")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[hr].height = 30

    r = hr
    for i, s in enumerate(S, 1):
        row = s["row"]
        plat = row["platform"]
        r += 1
        vals = [
            i, PLAT_CN.get(plat, plat), row.get("account_name"), row.get("account_followers"),
            row.get("video_url"), row.get("title"), row.get("cover_desc"),
            fmt(row.get("views")), fmt(row.get("likes")), fmt(row.get("comments")),
            fmt(row.get("shares")), fmt(row.get("collects")), fmt(row.get("danmaku")),
            (int(num(row.get("duration_sec"))) if num(row.get("duration_sec")) is not None else "N/A"),
            "爆款" if is_viral(row) else "对照", row.get("topic_category"), row.get("hook_type"),
            sample_reason_auto(s), s["reason"],
        ]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v)
            c.border = BORDER
            c.font = Font(size=9)
            c.alignment = Alignment(vertical="top", wrap_text=(j in (6, 7, 18, 19)))
            if j == 5 and v:
                c.hyperlink = v
                c.font = Font(size=9, color="0563C1", underline="single")
            if j in (4, 8, 9, 10, 11, 12, 13, 14):
                c.alignment = Alignment(horizontal="right", vertical="top")
            if j == 15:
                c.alignment = Alignment(horizontal="center", vertical="top")
                c.font = Font(size=9, bold=True, color="C00000" if v == "爆款" else "7F7F7F")
        if is_viral(row):
            ws.cell(row=r, column=15).fill = PatternFill("solid", fgColor="FCE4E4")
        else:
            ws.cell(row=r, column=15).fill = PatternFill("solid", fgColor="F2F2F2")

    for j, w in enumerate(WIDTH, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)

    # 按平台分布表
    ws2 = wb.create_sheet("选取分布")
    ws2.append(["平台", "样片数", "其中爆款", "其中对照"])
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5597")
    import collections
    cnt = collections.Counter((PLAT_CN.get(s["row"]["platform"]), is_viral(s["row"])) for s in S)
    for plat in ["B站", "TikTok", "抖音", "快手", "X (Twitter)"]:
        v = cnt.get((plat, True), 0)
        n = cnt.get((plat, False), 0)
        if v or n:
            ws2.append([plat, v + n, v, n])
    ws2.append([])
    ws2.append(["合计", len(S), sum(1 for s in S if is_viral(s["row"])), sum(1 for s in S if not is_viral(s["row"]))])
    ws2["A1"].alignment = Alignment(horizontal="center")
    for col, w in zip("ABCD", [14, 10, 10, 10]):
        ws2.column_dimensions[col].width = w
    for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, max_col=4):
        for c in row:
            c.border = BORDER
            if isinstance(c.value, (int, float)):
                c.alignment = Alignment(horizontal="right")

    wb.save(OUT)
    print(f"saved {OUT}")
    print(f"样片数：{len(S)}（爆款 {sum(1 for s in S if is_viral(s['row']))} / 对照 {sum(1 for s in S if not is_viral(s['row']))}）")
    for s in S[:3]:
        print("  示例:", s["row"]["account_name"], "|", sample_reason_auto(s)[:70])


if __name__ == "__main__":
    main()
