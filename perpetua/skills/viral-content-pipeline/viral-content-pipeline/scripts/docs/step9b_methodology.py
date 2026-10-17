# -*- coding: utf-8 -*-
"""渲染 爆款方法论 → Word(.docx) + PDF。内容模型见 kb_content.py。"""
import sys, os, datetime
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from kb_content import build_content, build_appendix
OUT_DIR = os.environ.get("PIPELINE_OUT", os.path.dirname(os.path.dirname(os.path.abspath(__file__))))


ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCX = os.path.join(OUT_DIR, "爆款方法论_20261014.docx")
PDF = os.path.join(OUT_DIR, "爆款方法论_20261014.pdf")
HTML = os.path.join(OUT_DIR, "_methodology.html")

CJK = "微软雅黑"
ACCENT = "1F3864"
ACCENT2 = "2F5597"
TODAY = "2026-10-14"


# ============================================================== Word
def render_docx(content):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(1.6)
    sec.left_margin = sec.right_margin = Cm(1.5)

    st = doc.styles["Normal"]
    st.font.name = CJK
    st.font.size = Pt(9.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.line_spacing = 1.25

    def shade(cell, color):
        el = OxmlElement("w:shd")
        el.set(qn("w:val"), "clear")
        el.set(qn("w:fill"), color)
        cell._tc.get_or_add_tcPr().append(el)

    def para(text, size=9.5, bold=False, color=None, before=0, after=3, align=None, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        if align:
            p.alignment = align
        r = p.add_run(text)
        r.font.size = Pt(size); r.bold = bold; r.italic = italic
        r.font.name = CJK; r._element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        return p

    def styled_table(headers, rows, widths=None):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        hdr = t.rows[0].cells
        for i, h in enumerate(headers):
            hdr[i].text = ""
            p = hdr[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(h))
            r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = RGBColor.from_string("FFFFFF")
            r.font.name = CJK; r._element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
            shade(hdr[i], ACCENT2)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                p = cells[i].paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                r = p.add_run(str(v))
                r.font.size = Pt(8)
                r.font.name = CJK; r._element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
                if ri % 2 == 1:
                    shade(cells[i], "F2F5FA")
            if widths:
                for i, w in enumerate(widths):
                    cells[i].width = Cm(w)
        if widths:
            for i, w in enumerate(widths):
                t.rows[0].cells[i].width = Cm(w)
        para("", size=4, after=0)
        return t

    # ---------- 抬头 ----------
    para("竞品爆款方法论", size=20, bold=True, color=ACCENT, after=2)
    para(f"基于 8 平台公开数据实测 · 采集日 {TODAY} · 版本 v1", size=10, color="595959", after=2)
    para("数据底座：competitor_viral_20261014.xlsx（288 条真实视频行，5 个平台可采）　配套：样片库_20261014.xlsx（41 条精选）",
         size=8.5, color="595959", after=2)
    para("用途：选题会 / 脚本评审 / 封面设计 / 投放前对表　｜　全部结论可逐条追溯到原始视频链接，无估算、无编造",
         size=8.5, color="595959", after=8)

    for blk in content:
        t = blk["type"]
        if t == "h1":
            para(blk["text"], size=13.5, bold=True, color=ACCENT, before=9, after=4)
        elif t == "h2":
            para(blk["text"], size=11, bold=True, color=ACCENT2, before=6, after=3)
        elif t == "p":
            if blk.get("style") == "note":
                para(blk["text"], size=8.5, color="595959", italic=True)
            else:
                para(blk["text"])
        elif t == "ul":
            for it in blk["items"]:
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.line_spacing = 1.2
                r = p.add_run(it)
                r.font.size = Pt(9)
                r.font.name = CJK; r._element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
        elif t == "callout":
            t2 = doc.add_table(rows=1, cols=1)
            t2.style = "Table Grid"
            c = t2.rows[0].cells[0]
            shade(c, "FFF2CC")
            c.text = ""
            p0 = c.paragraphs[0]
            p0.paragraph_format.space_after = Pt(2)
            r0 = p0.add_run(blk["title"])
            r0.bold = True; r0.font.size = Pt(9.5); r0.font.color.rgb = RGBColor.from_string("7F4F00")
            r0.font.name = CJK; r0._element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
            for it in blk["items"]:
                pp = c.add_paragraph()
                pp.paragraph_format.space_after = Pt(1)
                pp.paragraph_format.line_spacing = 1.15
                rr = pp.add_run("· " + it)
                rr.font.size = Pt(8.5)
                rr.font.name = CJK; rr._element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
            para("", size=4, after=0)
        elif t == "table":
            styled_table(blk["headers"], blk["rows"], blk.get("widths"))

    para("— 本手册由采集数据自动生成；每条结论对应 row 级可复核记录 —", size=8, color="808080",
         before=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.save(DOCX)
    print("saved", DOCX)


# ============================================================== HTML / PDF
def esc(s):
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def render_html(content):
    body = []
    for blk in content:
        t = blk["type"]
        if t == "h1":
            body.append(f'<h1>{esc(blk["text"])}</h1>')
        elif t == "h2":
            body.append(f'<h2>{esc(blk["text"])}</h2>')
        elif t == "p":
            cls = ' class="note"' if blk.get("style") == "note" else ""
            body.append(f'<p{cls}>{esc(blk["text"])}</p>')
        elif t == "ul":
            body.append("<ul>" + "".join(f"<li>{esc(i)}</li>" for i in blk["items"]) + "</ul>")
        elif t == "callout":
            body.append('<div class="callout"><div class="ct">' + esc(blk["title"]) + "</div><ul>" +
                        "".join(f"<li>{esc(i)}</li>" for i in blk["items"]) + "</ul></div>")
        elif t == "table":
            w = blk.get("widths")
            cols = ""
            if w:
                tot = sum(w)
                cols = "<colgroup>" + "".join(f'<col style="width:{x/tot*100:.2f}%">' for x in w) + "</colgroup>"
            th = "".join(f"<th>{esc(h)}</th>" for h in blk["headers"])
            trs = ""
            for row in blk["rows"]:
                trs += "<tr>" + "".join(f"<td>{esc(c)}</td>" for c in row) + "</tr>"
            body.append(f'<table>{cols}<thead><tr>{th}</tr></thead><tbody>{trs}</tbody></table>')
        elif t == "pagebreak":
            body.append('<div class="pb"></div>')

    html = f"""<!doctype html><html lang="zh-CN"><head><meta charset="utf-8">
<style>
@page {{ size: A4; margin: 14mm 12mm 14mm 12mm; }}
* {{ box-sizing: border-box; }}
body {{ font-family: "Noto Sans CJK SC","Noto Sans SC",sans-serif; font-size: 9.5pt; color: #1a1a1a; line-height: 1.5; margin: 0; }}
h1 {{ font-size: 14pt; color: #{ACCENT}; margin: 14px 0 6px; padding-bottom: 3px; border-bottom: 1.5px solid #{ACCENT2}; page-break-after: avoid; }}
h2 {{ font-size: 11pt; color: #{ACCENT2}; margin: 10px 0 4px; page-break-after: avoid; }}
p {{ margin: 0 0 5px; }}
p.note {{ font-size: 8.5pt; color: #666; font-style: italic; }}
ul {{ margin: 0 0 6px; padding-left: 16px; }}
li {{ margin-bottom: 2.5px; }}
table {{ width: 100%; border-collapse: collapse; margin: 5px 0 9px; page-break-inside: auto; }}
th {{ background: #{ACCENT2}; color: #fff; font-size: 8pt; padding: 3px 4px; border: 0.5pt solid #cfd8e6; text-align: left; }}
td {{ font-size: 8pt; padding: 3px 4px; border: 0.5pt solid #cfd8e6; vertical-align: top; }}
tbody tr:nth-child(even) td {{ background: #F2F5FA; }}
.callout {{ background: #FFF2CC; border: 0.75pt solid #E8C36A; border-radius: 3px; padding: 7px 9px; margin: 6px 0 9px; page-break-inside: avoid; }}
.callout .ct {{ font-weight: 700; color: #7F4F00; font-size: 9.5pt; margin-bottom: 3px; }}
.callout ul {{ margin: 0; }}
.callout li {{ font-size: 8.5pt; }}
.pb {{ page-break-after: always; }}
.cover {{ border-left: 4px solid #{ACCENT2}; padding-left: 10px; margin-bottom: 12px; }}
.cover .t {{ font-size: 22pt; font-weight: 700; color: #{ACCENT}; line-height: 1.2; }}
.cover .s {{ font-size: 10pt; color: #555; margin-top: 4px; }}
.cover .m {{ font-size: 8.5pt; color: #666; margin-top: 3px; }}
.foot {{ text-align: center; font-size: 7.5pt; color: #999; margin-top: 14px; }}
</style></head><body>
<div class="cover">
  <div class="t">竞品爆款方法论</div>
  <div class="s">基于 8 平台公开数据实测 · 采集日 {TODAY} · 版本 v1</div>
  <div class="m">数据底座：competitor_viral_20261014.xlsx（288 条真实视频行，5 个平台可采）　配套：样片库_20261014.xlsx（41 条精选样片）<br>
  用途：选题会 / 脚本评审 / 封面设计 / 投放前对表　｜　全部结论可逐条追溯到原始视频链接，无估算、无编造</div>
</div>
{''.join(body)}
<div class="foot">本手册由采集数据自动生成 · 每条结论对应 row 级可复核记录 · {TODAY}</div>
</body></html>"""
    open(HTML, "w", encoding="utf-8").write(html)
    return HTML


def render_pdf():
    from playwright.sync_api import sync_playwright
    with sync_playwright() as pw:
        b = pw.chromium.launch()
        pg = b.new_page()
        pg.goto("file://" + HTML, wait_until="load")
        pg.pdf(path=PDF, format="A4", print_background=True,
               margin={"top": "14mm", "bottom": "14mm", "left": "12mm", "right": "12mm"})
        b.close()
    print("saved", PDF)


if __name__ == "__main__":
    main_content = build_content()
    render_docx(main_content)
    render_html(main_content + build_appendix())
    render_pdf()
    # 页数
    try:
        from pypdf import PdfReader
        print("PDF 页数:", len(PdfReader(PDF).pages))
    except Exception as e:
        import re
        data = open(PDF, "rb").read()
        print("PDF 页数(正则):", len(re.findall(rb"/Type\s*/Page[^s]", data)))
